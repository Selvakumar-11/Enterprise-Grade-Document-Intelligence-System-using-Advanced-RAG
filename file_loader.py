# ============================================================
#  file_loader.py
#  Loads PDF, DOCX, and TXT files into LangChain Documents
#  with metadata: source, page_number, department
# ============================================================

from pathlib import Path
from langchain_core.documents import Document
from pypdf import PdfReader
from docx import Document as DocxDocument


def load_file(file_path: Path, department: str = "general") -> list[Document]:
    """
    Load a single file (PDF / DOCX / TXT) and return a list
    of LangChain Documents with metadata attached.
    """
    ext = file_path.suffix.lower()
    docs = []

    if ext == ".pdf":
        docs = _load_pdf(file_path, department)

    elif ext == ".docx":
        docs = _load_docx(file_path, department)

    elif ext == ".txt":
        docs = _load_txt(file_path, department)

    else:
        print(f"⚠️  Unsupported file type: {ext} — skipping {file_path.name}")

    return docs


def load_files(file_list: list[tuple[str, str]]) -> list[Document]:
    """
    Load multiple files at once.

    Args:
        file_list: List of (file_path_str, department) tuples
                   e.g. [("contract.pdf", "legal"), ("policy.docx", "hr")]

    Returns:
        Combined list of Documents from all files.
    """
    all_docs = []
    for file_name, department in file_list:
        file_path = Path(file_name)
        if not file_path.exists():
            print(f"⚠️  File not found: {file_name} — skipping.")
            continue
        print(f"📄 Loading: {file_name} [{department}]")
        all_docs.extend(load_file(file_path, department))

    print(f"✅ Total raw documents loaded: {len(all_docs)}")
    return all_docs


# ── Private helpers ──────────────────────────────────────────

def _build_metadata(source: str, page_number: int, department: str) -> dict:
    return {
        "source": source,
        "page_number": page_number,
        "department": department,
    }


def _load_pdf(file_path: Path, department: str) -> list[Document]:
    reader = PdfReader(str(file_path))
    docs = []
    for i, page in enumerate(reader.pages):
        text = page.extract_text() or ""
        if text.strip():
            docs.append(Document(
                page_content=text,
                metadata=_build_metadata(file_path.name, i + 1, department),
            ))
    return docs


def _load_docx(file_path: Path, department: str) -> list[Document]:
    docx = DocxDocument(str(file_path))
    full_text = "\n".join([p.text for p in docx.paragraphs if p.text.strip()])
    return [Document(
        page_content=full_text,
        metadata=_build_metadata(file_path.name, 1, department),
    )]


def _load_txt(file_path: Path, department: str) -> list[Document]:
    with open(file_path, "r", encoding="utf-8") as f:
        text = f.read()
    return [Document(
        page_content=text,
        metadata=_build_metadata(file_path.name, 1, department),
    )]